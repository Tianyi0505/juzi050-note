# -*- coding: utf-8 -*-
"""把 设计图/ 下的 PNG 合成 Word 文档（课程设计要求：图片复制到 docx）。

用法（任意目录执行都可以）：
    python 设计图/make_docx.py
输出写到 设计图/设计图集.docx。
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

# 项目根目录：校园快递收发系统
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "设计图"

FIGURES = [
    ("1-业务流程图.png", "图 1  校园快递收发业务流程图",
     "说明：按角色标注流程节点。蓝色为学生，绿色为工作人员，紫色为系统。"),
    ("2-模块结构图.png", "图 2  系统模块结构图",
     "说明：后端按业务拆分为 6 个模块。"),
    ("3-程序流程图-包裹入库.png", "图 3  程序流程图：包裹入库",
     "说明：对应 internal/handler/parcel.go 的 CreateParcel 函数。"),
    ("4-程序流程图-取件核销.png", "图 4  程序流程图：取件核销",
     "说明：对应 internal/handler/parcel.go 的 doPickup 函数。"),
    ("5-E-R图.png", "图 5  数据库 E-R 图",
     "说明：5 张表。PK 为主键，FK 为外键。"),
]

MAX_WIDTH_CM = 16.0
MAX_HEIGHT_CM = 20.5


def main():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("校园快递收发系统  设计图集")
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(20)
    run.bold = True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("软件工程课程设计").font.size = Pt(12)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(12)
    note.add_run("本文件包含 5 张设计图：业务流程图、模块结构图、两张程序流程图、数据库 E-R 图。"
                 "图片由 设计图/gen_diagrams.py 生成，可重新运行脚本复现。")

    for filename, caption, desc in FIGURES:
        path = OUT / filename
        if not path.exists():
            raise SystemExit(f"缺少图片：{path}，请先运行 gen_diagrams.py")

        doc.add_page_break()

        h = doc.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hr = h.add_run(caption)
        hr.font.name = "黑体"
        hr._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        hr.font.size = Pt(14)
        hr.bold = True

        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        w_px, h_px = Image.open(path).size
        if MAX_WIDTH_CM * h_px / w_px > MAX_HEIGHT_CM:
            pic.add_run().add_picture(str(path), height=Cm(MAX_HEIGHT_CM))
        else:
            pic.add_run().add_picture(str(path), width=Cm(MAX_WIDTH_CM))

        d = doc.add_paragraph()
        d.alignment = WD_ALIGN_PARAGRAPH.CENTER
        d.add_run(desc).font.size = Pt(9)

    out = OUT / "设计图集.docx"
    doc.save(str(out))
    print("saved", out.relative_to(ROOT))


if __name__ == "__main__":
    main()
